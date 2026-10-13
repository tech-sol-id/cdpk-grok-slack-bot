import os
import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

def generate_bom_excel(project_name: str, client_name: str = "", items: list = None, output_dir: str = "/tmp") -> str:
    """Menghasilkan spreadsheet Microsoft Excel (.xlsx) Bill of Materials (BOM) resmi CDPK HQ."""
    os.makedirs(output_dir, exist_ok=True)
    items = items or []
    safe_pname = "".join(c for c in project_name if c.isalnum() or c in (" ", "_", "-")).strip().replace(" ", "_")
    now_str = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    file_path = os.path.join(output_dir, f"BOM_CDPK_{safe_pname}_{now_str}.xlsx")

    wb = Workbook()
    ws = wb.active
    ws.title = "Bill of Materials"
    ws.views.sheetView[0].showGridLines = True

    # Brand Colors (CDPK Crimson & Slate)
    c_crimson = "C8102E"
    c_dark_slate = "1E232A"
    c_light_bg = "F4F5F7"

    font_title = Font(name="Segoe UI", size=16, bold=True, color="FFFFFF")
    font_subtitle = Font(name="Segoe UI", size=10, italic=True, color="FFFFFF")
    font_header = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    font_data = Font(name="Segoe UI", size=10, color="000000")
    font_bold = Font(name="Segoe UI", size=10, bold=True, color="000000")

    fill_title = PatternFill(start_color=c_crimson, end_color=c_crimson, fill_type="solid")
    fill_header = PatternFill(start_color=c_dark_slate, end_color=c_dark_slate, fill_type="solid")
    fill_zebra = PatternFill(start_color=c_light_bg, end_color=c_light_bg, fill_type="solid")
    fill_total = PatternFill(start_color="FFE8EC", end_color="FFE8EC", fill_type="solid")

    thin_border = Border(
        left=Side(style='thin', color='DDDDDD'),
        right=Side(style='thin', color='DDDDDD'),
        top=Side(style='thin', color='DDDDDD'),
        bottom=Side(style='thin', color='DDDDDD')
    )
    thick_bottom = Border(bottom=Side(style='double', color='000000'), top=Side(style='thin', color='AAAAAA'))

    # Title Banner
    ws.merge_cells("A1:H1")
    ws["A1"] = f"CDPK HQ — BILL OF MATERIALS (BOM)"
    ws["A1"].font = font_title
    ws["A1"].fill = fill_title
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 36

    ws.merge_cells("A2:H2")
    ws["A2"] = f"PT Citradata Purnakharisma & Techsolutions | Proyek: {project_name} | Klien: {client_name or 'Umum'} | Tanggal: {datetime.date.today().strftime('%d %B %Y')}"
    ws["A2"].font = font_subtitle
    ws["A2"].fill = PatternFill(start_color="9B0D23", end_color="9B0D23", fill_type="solid")
    ws["A2"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 20

    ws.append([]) # row 3 blank

    # Table Headers
    headers = [
        "No", "Part Number / Model", "Deskripsi Perangkat", "Brand",
        "Qty", "Satuan", "Est. Harga Satuan (IDR)", "Total Harga (IDR)"
    ]
    ws.append(headers)
    ws.row_dimensions[4].height = 26
    for col_idx in range(1, 9):
        cell = ws.cell(row=4, column=col_idx)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # Data Rows
    start_row = 5
    for idx, itm in enumerate(items, 1):
        r = start_row + idx - 1
        pnum = itm.get("part_number") or itm.get("model") or "-"
        desc = itm.get("description") or "-"
        brand = itm.get("brand") or ("FMTCS" if "omu" in desc.lower() or "oru" in desc.lower() or "coupler" in desc.lower() else ("Motorola" if "repeater" in desc.lower() or "slr" in desc.lower() or "radio" in desc.lower() else "CDPK"))
        qty = int(itm.get("quantity") or itm.get("qty") or 1)
        unit = itm.get("unit") or "Unit"
        unit_price = float(itm.get("unit_price") or itm.get("price") or 0.0)

        ws.cell(row=r, column=1, value=idx).alignment = Alignment(horizontal="center")
        ws.cell(row=r, column=2, value=pnum).alignment = Alignment(horizontal="center")
        ws.cell(row=r, column=3, value=desc).alignment = Alignment(horizontal="left")
        ws.cell(row=r, column=4, value=brand).alignment = Alignment(horizontal="center")
        ws.cell(row=r, column=5, value=qty).alignment = Alignment(horizontal="center")
        ws.cell(row=r, column=6, value=unit).alignment = Alignment(horizontal="center")

        c_uprice = ws.cell(row=r, column=7, value=unit_price)
        c_uprice.alignment = Alignment(horizontal="right")
        c_uprice.number_format = '"Rp "#,##0'

        # Live formula Total Harga = Qty * Unit Price
        c_tot = ws.cell(row=r, column=8, value=f"=E{r}*G{r}")
        c_tot.alignment = Alignment(horizontal="right")
        c_tot.number_format = '"Rp "#,##0'

        for col_idx in range(1, 9):
            c = ws.cell(row=r, column=col_idx)
            c.font = font_data
            c.border = thin_border
            if idx % 2 == 0:
                c.fill = fill_zebra
        ws.row_dimensions[r].height = 22

    end_row = start_row + len(items) - 1 if items else start_row
    subtotal_row = end_row + 1
    ppn_row = subtotal_row + 1
    grand_row = ppn_row + 1

    # Subtotal
    ws.merge_cells(f"A{subtotal_row}:G{subtotal_row}")
    ws[f"A{subtotal_row}"] = "SUBTOTAL"
    ws[f"A{subtotal_row}"].font = font_bold
    ws[f"A{subtotal_row}"].alignment = Alignment(horizontal="right")
    ws[f"H{subtotal_row}"] = f"=SUM(H{start_row}:H{end_row})" if items else 0
    ws[f"H{subtotal_row}"].font = font_bold
    ws[f"H{subtotal_row}"].number_format = '"Rp "#,##0'

    # PPN 11%
    ws.merge_cells(f"A{ppn_row}:G{ppn_row}")
    ws[f"A{ppn_row}"] = "PPN (11%)"
    ws[f"A{ppn_row}"].font = font_bold
    ws[f"A{ppn_row}"].alignment = Alignment(horizontal="right")
    ws[f"H{ppn_row}"] = f"=H{subtotal_row}*0.11"
    ws[f"H{ppn_row}"].font = font_bold
    ws[f"H{ppn_row}"].number_format = '"Rp "#,##0'

    # Grand Total
    ws.merge_cells(f"A{grand_row}:G{grand_row}")
    ws[f"A{grand_row}"] = "GRAND TOTAL"
    ws[f"A{grand_row}"].font = Font(name="Segoe UI", size=11, bold=True, color=c_crimson)
    ws[f"A{grand_row}"].alignment = Alignment(horizontal="right")
    ws[f"A{grand_row}"].fill = fill_total
    ws[f"H{grand_row}"] = f"=H{subtotal_row}+H{ppn_row}"
    ws[f"H{grand_row}"].font = Font(name="Segoe UI", size=11, bold=True, color=c_crimson)
    ws[f"H{grand_row}"].alignment = Alignment(horizontal="right")
    ws[f"H{grand_row}"].number_format = '"Rp "#,##0'
    ws[f"H{grand_row}"].fill = fill_total
    ws[f"H{grand_row}"].border = thick_bottom

    # Auto-adjust column width
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)
    ws.column_dimensions["C"].width = 42 # Deskripsi wider

    wb.save(file_path)
    return file_path

def generate_proposal_docx(
    project_name: str,
    client_name: str = "",
    scope_summary: str = "",
    topology_description: str = "",
    items: list = None,
    technical_notes: str = "",
    output_dir: str = "/tmp"
) -> str:
    """Menghasilkan Proposal Solusi Teknis & Arsitektur (.docx) resmi CDPK HQ."""
    os.makedirs(output_dir, exist_ok=True)
    safe_pname = "".join(c for c in project_name if c.isalnum() or c in (" ", "_", "-")).strip().replace(" ", "_")
    now_str = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    file_path = os.path.join(output_dir, f"Proposal_Teknis_CDPK_{safe_pname}_{now_str}.docx")

    doc = docx.Document()

    # Set Margins 1 inch
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Styles
    p_header = doc.add_paragraph()
    p_header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_corp = p_header.add_run("PT CITRADATA PURNAKHARISMA & TECHSOLUTIONS\n")
    r_corp.font.name = "Calibri"
    r_corp.font.size = Pt(9)
    r_corp.font.bold = True
    r_corp.font.color.rgb = RGBColor(200, 16, 46) # Crimson Red
    r_subcorp = p_header.add_run("Authorized Motorola Solutions & FMTCS Solutions Partner Indonesia\nJl. Biak No. 17, Gambir, Jakarta Pusat 10150")
    r_subcorp.font.name = "Calibri"
    r_subcorp.font.size = Pt(8)
    r_subcorp.font.color.rgb = RGBColor(120, 120, 120)

    # Title
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(18)
    p_title.paragraph_format.space_after = Pt(6)
    r_title = p_title.add_run(f"PROPOSAL DESAIN SOLUSI & ARSITEKTUR TEKNIS")
    r_title.font.name = "Calibri"
    r_title.font.size = Pt(18)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(30, 35, 42)

    p_proj = doc.add_paragraph()
    p_proj.paragraph_format.space_after = Pt(20)
    r_proj = p_proj.add_run(f"Proyek: {project_name}\nKlien: {client_name or 'Mitra Kerja CDPK'}\nTanggal: {datetime.date.today().strftime('%d %B %Y')}")
    r_proj.font.name = "Calibri"
    r_proj.font.size = Pt(11)
    r_proj.font.color.rgb = RGBColor(80, 80, 80)

    # 1. Executive Summary
    h1 = doc.add_heading("1. Ringkasan Eksekutif & Latar Belakang", level=1)
    h1.style.font.color.rgb = RGBColor(200, 16, 46)
    p_exec = doc.add_paragraph(scope_summary or f"Proposal ini memuat rancangan arsitektur sistem telekomunikasi nirkabel profesional untuk {project_name}, guna memastikan ketersediaan cakupan sinyal (100% zero blind-spot) pada seluruh area kritis gedung dan fasilitas operasional.")
    p_exec.style.font.name = "Calibri"
    p_exec.style.font.size = Pt(10.5)

    # 2. Arsitektur & Topologi
    h2 = doc.add_heading("2. Analisis Topologi & Rekayasa Sistem (DAS)", level=1)
    h2.style.font.color.rgb = RGBColor(200, 16, 46)
    p_topo = doc.add_paragraph(topology_description or "Sistem dirancang mengadopsi teknologi Active Optical Distributed Antenna System (DAS) FMTCS dipadukan dengan repeater digital Motorola MOTOTRBO. Penggunaan fiber optik antar Optical Master Unit (OMU) dan Optical Remote Unit (ORU) mengeliminasi redaman sinyal pada bentang kabel vertikal yang panjang, sementara jaringan pasif coaxial coupler dan indoor omni antenna mendistribusikan sinyal secara presisi di setiap lantai.")
    p_topo.style.font.name = "Calibri"
    p_topo.style.font.size = Pt(10.5)

    # 3. Rekomendasi Perangkat Utama
    h3 = doc.add_heading("3. Usulan Perangkat & Bill of Materials (BOM)", level=1)
    h3.style.font.color.rgb = RGBColor(200, 16, 46)

    if items:
        table = doc.add_table(rows=1, cols=5)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.style = 'Table Grid'
        hdr_cells = table.rows[0].cells
        headers = ["No", "Part Number / Model", "Deskripsi Perangkat", "Brand", "Kuantitas"]
        for i, h in enumerate(headers):
            hdr_cells[i].text = h
            hdr_cells[i].paragraphs[0].runs[0].font.bold = True
            hdr_cells[i].paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
            shading = hdr_cells[i]._tc.get_or_add_tcPr()
            from docx.oxml import OxmlElement
            from docx.oxml.ns import qn
            shd = OxmlElement('w:shd')
            shd.set(qn('w:val'), 'clear')
            shd.set(qn('w:color'), 'auto')
            shd.set(qn('w:fill'), 'C8102E')
            shading.append(shd)

        for idx, itm in enumerate(items, 1):
            row_cells = table.add_row().cells
            row_cells[0].text = str(idx)
            row_cells[1].text = itm.get("part_number") or itm.get("model") or "-"
            row_cells[2].text = itm.get("description") or "-"
            row_cells[3].text = itm.get("brand") or "CDPK/FMTCS"
            row_cells[4].text = f"{itm.get('quantity') or 1} {itm.get('unit') or 'Unit'}"
    else:
        doc.add_paragraph("Daftar Bill of Materials terlampir pada dokumen perhitungan terpisah.")

    # 4. Layanan Purna Jual & Garansi
    h4 = doc.add_heading("4. Jaminan Kualitas & Dukungan Resmi CDPK HQ", level=1)
    h4.style.font.color.rgb = RGBColor(200, 16, 46)
    p_sla = doc.add_paragraph(
        "• Pengalaman Teruji: PT Citradata Purnakharisma melayani solusi komunikasi radio nirkabel sejak 1993 (30+ tahun distributor resmi Motorola Solutions).\n"
        "• Pusat Servis Resmi (Authorized Service Center): Didukung teknisi bersertifikasi dan peralatan kalibrasi laboratorium radio terlengkap.\n"
        "• Sertifikasi Terakreditasi: Sistem manajemen mutu ISO 9001:2015, K3 ISO 45001:2018, dan Lingkungan ISO 14001:2015 menjamin integritas pelaksanaan proyek."
    )
    p_sla.style.font.name = "Calibri"
    p_sla.style.font.size = Pt(10)

    if technical_notes:
        h5 = doc.add_heading("5. Catatan Teknis & Rekomendasi Pelaksanaan", level=1)
        h5.style.font.color.rgb = RGBColor(200, 16, 46)
        doc.add_paragraph(technical_notes)

    doc.save(file_path)
    return file_path

def upload_document_to_slack(client, channel: str, thread_ts: str, file_path: str, title: str, initial_comment: str = "") -> bool:
    """Mengunggah dokumen (.xlsx, .docx, .pdf) langsung ke thread Slack menggunakan Slack API files.upload_v2."""
    if not os.path.exists(file_path):
        return False

    try:
        fname = os.path.basename(file_path)
        client.files_upload_v2(
            channel=channel,
            thread_ts=thread_ts,
            file=file_path,
            filename=fname,
            title=title or fname,
            initial_comment=initial_comment
        )
        return True
    except Exception as e:
        print(f"Error uploading file to Slack: {e}")
        return False

DOC_TOOLS = [
    {
        "type": "function",
        "function": {
            "name": "generate_bom_excel",
            "description": "Menghasilkan spreadsheet Microsoft Excel (.xlsx) resmi CDPK HQ berisi Bill of Materials (BOM), kuantitas, part number, dan estimasi harga perangkat (Motorola, FMTCS, i.Safe). Panggil fungsi ini jika pengguna meminta file Excel, spreadsheet BOM, atau rincian daftar harga/BOM.",
            "parameters": {
                "type": "object",
                "properties": {
                    "project_name": {"type": "string", "description": "Nama proyek atau nama gedung"},
                    "client_name": {"type": "string", "description": "Nama perusahaan klien atau mitra"},
                    "items": {
                        "type": "array",
                        "description": "Daftar perangkat dalam BOM",
                        "items": {
                            "type": "object",
                            "properties": {
                                "part_number": {"type": "string", "description": "Model atau Part Number alat (misal: SLR5000, LRFN-DOMU400C)"},
                                "description": {"type": "string", "description": "Nama atau deskripsi lengkap perangkat"},
                                "brand": {"type": "string", "description": "Brand perangkat (Motorola, FMTCS, i.Safe)"},
                                "quantity": {"type": "integer", "description": "Jumlah kuantitas unit"},
                                "unit": {"type": "string", "description": "Satuan (Unit, Pcs, Roll, Lot)"},
                                "unit_price": {"type": "number", "description": "Estimasi harga satuan dalam IDR"}
                            },
                            "required": ["description", "quantity"]
                        }
                    }
                },
                "required": ["project_name", "items"]
            }
        }
    },
    {
        "type": "function",
        "function": {
            "name": "generate_proposal_docx",
            "description": "Menghasilkan dokumen resmi Microsoft Word (.docx) Proposal Desain Solusi Teknis & Arsitektur CDPK HQ. Panggil fungsi ini jika pengguna meminta proposal tertulis, dokumen solusi teknis, dokumen Word, atau dokumen rancangan topologi proyek.",
            "parameters": {
                "type": "object",
                "properties": {
                    "project_name": {"type": "string", "description": "Nama proyek atau nama gedung"},
                    "client_name": {"type": "string", "description": "Nama perusahaan klien"},
                    "scope_summary": {"type": "string", "description": "Latar belakang dan ringkasan cakupan proyek"},
                    "topology_description": {"type": "string", "description": "Analisis dan penjelasan rekayasa arsitektur topologi jaringan DAS / Radio"},
                    "items": {
                        "type": "array",
                        "description": "Daftar perangkat utama dalam proposal",
                        "items": {
                            "type": "object",
                            "properties": {
                                "part_number": {"type": "string"},
                                "description": {"type": "string"},
                                "brand": {"type": "string"},
                                "quantity": {"type": "integer"},
                                "unit": {"type": "string"}
                            },
                            "required": ["description", "quantity"]
                        }
                    },
                    "technical_notes": {"type": "string", "description": "Catatan teknis, instalasi, dan kepatuhan standar"}
                },
                "required": ["project_name", "scope_summary", "topology_description"]
            }
        }
    }
]

